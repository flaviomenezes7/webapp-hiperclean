"""Examina a estrutura da planilha do Helsio para entender o formato."""
import openpyxl
import sys

wb = openpyxl.load_workbook(r"C:\Users\mnzfl\Desktop\Helsio Excel.xlsx", data_only=True)

print(f"Abas encontradas: {wb.sheetnames}")
print()

for sheet_name in wb.sheetnames[:2]:  # Examinar as 2 primeiras abas
    ws = wb[sheet_name]
    print(f"=== Aba: {sheet_name} ===")
    print(f"Linhas: {ws.max_row}, Colunas: {ws.max_column}")
    print()

    # Print first 25 rows to understand the structure
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(25, ws.max_row), values_only=False), 1):
        values = []
        for cell in row[:8]:  # Columns A-H
            v = cell.value
            if v is not None:
                v_str = str(v)[:50]
            else:
                v_str = ""
            values.append(f"{cell.column_letter}={v_str}")
        print(f"  Row {row_idx}: {' | '.join(values)}")
    print()

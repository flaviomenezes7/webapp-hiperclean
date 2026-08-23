"""
Servico de importacao de planilha Excel do Helsio.
Logica de parse e deduplicacao extraida do script import_excel.py.
"""

import re
import io
from datetime import date, datetime

import openpyxl
from sqlmodel import Session, select

from app.models.cliente import Cliente
from app.models.atendimento import Atendimento
from app.models.envio import Envio
from app.models.enums import TipoServico


def clean_nome(raw: str) -> str | None:
    if not raw or not raw.strip():
        return None

    text = raw.strip()

    skip_patterns = [
        r'^F\s*O\s*L\s*G\s*A',
        r'^s[aá]bado',
        r'^domingo',
        r'^Viagem',
        r'^Manuten[cç][aã]o',
        r'^Cliente$',
        r'^Janeiro$', r'^Fevereiro$', r'^Mar[cç]o$', r'^Abril$', r'^Maio$',
        r'^Junho$', r'^Julho$', r'^Agosto$', r'^Setembro$', r'^Outubro$',
        r'^Novembro$', r'^Dezembro$',
        r'^Agenda de Clientes',
        r'^Oficina',
        r'^Helzio',
        r'^n[aã]o fez',
    ]
    for pattern in skip_patterns:
        if re.search(pattern, text, re.IGNORECASE):
            return None

    text = re.sub(r'\s+\d{1,2}h\d{0,2}\s*', ' ', text)
    text = re.sub(r'\s+\d{1,2}\.\d{3}[oa\u00ba\u00aa]?\s*(cliente|clien|cl)\w*', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s+\d{3,4}[oa\u00ba\u00aa]?\s*(cliente|clien|cl)\w*', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s*\((?:Reagendamos|desmarcou|WhatsApp|mae\s|filha\s|o carpete|nao fez)[^)]*\)', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s+desmarcou.*$', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s+', ' ', text).strip()

    if len(text) < 2:
        return None

    return text


def clean_telefone(raw) -> str | None:
    if not raw:
        return None

    text = str(raw).strip()
    if not text:
        return None

    parts = re.split(r'\s{2,}', text)
    text = parts[0].strip()
    digits = re.sub(r'\D', '', text)

    if len(digits) < 8:
        return None

    if len(digits) >= 12 and digits.startswith('55'):
        return digits[:13]

    if len(digits) >= 10:
        return '55' + digits[:11]

    if 8 <= len(digits) <= 9:
        return digits

    return None


def parse_aniversario(raw) -> date | None:
    if raw is None:
        return None

    if isinstance(raw, datetime):
        return date(raw.year, raw.month, raw.day)

    text = str(raw).strip()
    if not text:
        return None

    meses = {
        'jan': 1, 'fev': 2, 'mar': 3, 'abr': 4, 'mai': 5, 'jun': 6,
        'jul': 7, 'ago': 8, 'set': 9, 'out': 10, 'nov': 11, 'dez': 12,
    }
    match = re.match(r'(\d{1,2})[/-](\w{3})', text, re.IGNORECASE)
    if match:
        day = int(match.group(1))
        month_str = match.group(2).lower()
        month = meses.get(month_str)
        if month and 1 <= day <= 31:
            return date(2000, month, day)

    return None


def classify_servico(descricao: str) -> TipoServico:
    if not descricao:
        return TipoServico.OUTRO

    text = descricao.lower()

    if 'colch' in text:
        return TipoServico.COLCHAO
    if 'cadeira' in text:
        return TipoServico.CADEIRA
    if 'tapete' in text or 'carpete' in text:
        return TipoServico.TAPETE
    if 'cortina' in text:
        return TipoServico.CORTINA
    if 'poltrona' in text:
        return TipoServico.POLTRONA
    if 'banco' in text and ('carro' in text or 'veic' in text or 'automotivo' in text):
        return TipoServico.BANCO_AUTOMOTIVO
    if 'sof' in text:
        return TipoServico.SOFA
    if 'puff' in text or 'puf' in text:
        return TipoServico.SOFA

    return TipoServico.OUTRO


def parse_date(raw) -> date | None:
    if raw is None:
        return None

    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw

    text = str(raw).strip()
    match = re.search(r'(\d{2})/(\d{2})/(\d{4})', text)
    if match:
        try:
            return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            pass

    return None


def importar_excel(session: Session, file_bytes: bytes, limpar_dados: bool = True) -> dict:
    """
    Importa dados de uma planilha Excel para o banco.
    
    Args:
        session: Database session
        file_bytes: Conteudo do arquivo Excel em bytes
        limpar_dados: Se True, remove todos os dados existentes antes de importar
    
    Returns:
        Dict com estatisticas da importacao
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    # Collect all raw records
    records = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            col_a = row[0] if len(row) > 0 else None
            col_b = row[1] if len(row) > 1 else None
            col_c = row[2] if len(row) > 2 else None
            col_e = row[4] if len(row) > 4 else None
            col_f = row[5] if len(row) > 5 else None
            col_g = row[6] if len(row) > 6 else None

            nome = clean_nome(str(col_a) if col_a else "")
            if not nome:
                continue

            telefone = clean_telefone(col_e)
            aniversario = parse_aniversario(col_f)
            endereco = str(col_g).strip() if col_g else None
            data_atend = parse_date(col_c)
            tipo_servico = classify_servico(str(col_b) if col_b else "")
            descricao_servico = str(col_b).strip() if col_b else None

            records.append({
                'nome': nome,
                'telefone': telefone,
                'aniversario': aniversario,
                'endereco': endereco,
                'data_atend': data_atend,
                'tipo_servico': tipo_servico,
                'descricao_servico': descricao_servico,
            })

    # Deduplicate by phone or name
    clientes_map = {}

    for rec in records:
        if rec['telefone'] and len(rec['telefone']) >= 8:
            key = rec['telefone'][-8:]
        else:
            key = rec['nome'].lower().strip()

        if key not in clientes_map:
            clientes_map[key] = {
                'nome': rec['nome'],
                'telefone': rec['telefone'],
                'aniversario': rec['aniversario'],
                'endereco': rec['endereco'],
                'atendimentos': [],
            }
        else:
            existing = clientes_map[key]
            if not existing['telefone'] and rec['telefone']:
                existing['telefone'] = rec['telefone']
            if not existing['aniversario'] and rec['aniversario']:
                existing['aniversario'] = rec['aniversario']
            if not existing['endereco'] and rec['endereco']:
                existing['endereco'] = rec['endereco']

        if rec['data_atend']:
            clientes_map[key]['atendimentos'].append({
                'data_atend': rec['data_atend'],
                'tipo_servico': rec['tipo_servico'],
                'descricao_servico': rec['descricao_servico'],
            })

    # Remove entries without phone
    clientes_map = {k: v for k, v in clientes_map.items() if v['telefone']}

    # Clear existing data if requested
    if limpar_dados:
        envios = session.exec(select(Envio)).all()
        for e in envios:
            session.delete(e)

        atends_existing = session.exec(select(Atendimento)).all()
        for a in atends_existing:
            session.delete(a)

        existing_clients = session.exec(select(Cliente)).all()
        for c in existing_clients:
            session.delete(c)

        session.commit()

    # Insert new data
    clientes_inseridos = 0
    atendimentos_inseridos = 0

    for key, data in clientes_map.items():
        cliente = Cliente(
            nome=data['nome'],
            telefone=data['telefone'] or "",
            data_nasc=data['aniversario'],
            endereco=data['endereco'],
        )
        session.add(cliente)
        session.flush()
        clientes_inseridos += 1

        for atend in data['atendimentos']:
            atendimento = Atendimento(
                cliente_id=cliente.id,
                tipo_servico=atend['tipo_servico'],
                data_atend=atend['data_atend'],
                observacoes=atend['descricao_servico'],
            )
            session.add(atendimento)
            atendimentos_inseridos += 1

    session.commit()

    return {
        "registros_brutos": len(records),
        "clientes_inseridos": clientes_inseridos,
        "atendimentos_inseridos": atendimentos_inseridos,
        "abas_processadas": wb.sheetnames,
    }
